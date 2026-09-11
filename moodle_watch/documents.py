"""Rendre le contenu d'un fichier lisible par un agent.

Deux chemins, dans cet ordre. Si le miroir a deja le fichier sur le disque, on le
lit : c'est gratuit et ca marche hors ligne. Sinon on le telecharge par le jeton.

Le second chemin repose sur un point verifie plutot que suppose :
`webservice/pluginfile.php` accepte le jeton du service web en parametre de
requete. C'est ce qui permet de lire un document sans avoir a lancer une
aspiration complete.
"""

from __future__ import annotations

import re
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

#: Au-dela, on rend les metadonnees et pas le contenu. Une video ne se resume pas
#: ici, et charger 6 Go dans une reponse d'outil n'aide personne.
TAILLE_MAX_TELECHARGEMENT = 40 * 1024 * 1024

EXT_TEXTE = {".md", ".txt", ".csv", ".json", ".srt", ".vtt", ".url", ".webloc"}
EXT_MEDIA = {".mp4", ".mkv", ".webm", ".mov", ".avi", ".mp3", ".wav", ".m4a", ".ogg",
             ".zip", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"}

_BALISE = re.compile(r"<[^>]+>")
_BLANCS = re.compile(r"\n{3,}")


def _pages(spec: str | None, total: int) -> range:
    if not spec:
        return range(total)
    m = re.match(r"^\s*(\d+)\s*(?:-\s*(\d+))?\s*$", spec)
    if not m:
        raise ValueError(f"`pages` incomprehensible : {spec!r}. Exemples : 3, 1-5.")
    debut = max(1, int(m.group(1)))
    fin = int(m.group(2) or debut)
    return range(debut - 1, min(fin, total))


def url_avec_jeton(fileurl: str, token: str) -> str:
    """Ajoute le jeton a une URL de `pluginfile.php`.

    Le jeton part en parametre de requete parce que c'est la seule forme que
    `pluginfile.php` accepte. L'URL ne doit donc jamais etre journalisee ni
    rendue a un agent : elle porte une capacite complete sur le compte.
    """
    morceaux = urllib.parse.urlsplit(fileurl)
    params = urllib.parse.parse_qsl(morceaux.query, keep_blank_values=True)
    params = [(k, v) for k, v in params if k != "token"]
    params.append(("token", token))
    return urllib.parse.urlunsplit(morceaux._replace(query=urllib.parse.urlencode(params)))


def telecharger(fileurl: str, token: str, *, limite: int = TAILLE_MAX_TELECHARGEMENT) -> bytes:
    requete = urllib.request.Request(
        url_avec_jeton(fileurl, token), headers={"User-Agent": "moodle-watch"}
    )
    with urllib.request.urlopen(requete, timeout=60) as rep:
        donnees = rep.read(limite + 1)
    if len(donnees) > limite:
        raise ValueError(f"fichier plus gros que la limite de {limite} octets")
    # Moodle rend 200 avec un corps JSON quand le jeton est refuse, pas un 403.
    if donnees[:1] == b"{" and b"errorcode" in donnees[:400]:
        raise PermissionError(donnees[:400].decode("utf-8", "replace"))
    return donnees


def _texte_pdf(donnees: bytes, spec: str | None) -> tuple[str, int]:
    import pymupdf

    with pymupdf.open(stream=donnees, filetype="pdf") as doc:
        indices = _pages(spec, doc.page_count)
        return "\n\n".join(doc[i].get_text() for i in indices), doc.page_count


def _texte_docx(donnees: bytes) -> str:
    import io

    from docx import Document

    doc = Document(io.BytesIO(donnees))
    morceaux = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for ligne in table.rows:
            morceaux.append(" | ".join(c.text.strip() for c in ligne.cells))
    return "\n".join(morceaux)


def _texte_html(donnees: bytes) -> str:
    brut = donnees.decode("utf-8", "replace")
    brut = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", brut)
    return _BLANCS.sub("\n\n", _BALISE.sub(" ", brut)).strip()


def extract_text(row: dict[str, Any], engine: Any, *, pages: str | None = None,
                 max_chars: int = 40000) -> dict[str, Any]:
    """Rend le texte d'un fichier, ou ses metadonnees s'il n'en a pas."""
    nom = row["content_filename"]
    ext = Path(nom).suffix.lower()
    meta = {
        "file_id": row["file_id"],
        "filename": nom,
        "modname": row["module_modname"],
        "size": row["content_filesize"],
        "course_id": row["course_id"],
        "module": row["module_name"],
    }

    if ext in EXT_MEDIA:
        return meta | {"status": "not_text", "detail": f"{ext} is a media or archive file; no text extracted."}

    local = Path(row["saved_to"] or "")
    source = "local"
    if local.is_file():
        donnees = local.read_bytes()
    else:
        source = "download"
        if not row["content_fileurl"]:
            return meta | {"status": "no_source", "detail": "No local copy and no file URL."}
        jeton = engine.raw_config().get("token")
        if not jeton:
            return meta | {"status": "no_token"}
        try:
            donnees = telecharger(row["content_fileurl"], jeton)
        except PermissionError as err:
            from .engine import TokenExpired

            raise TokenExpired(str(err)) from err
        except Exception as err:  # noqa: BLE001
            return meta | {"status": "download_failed", "detail": f"{type(err).__name__}: {err}"}

    try:
        if ext == ".pdf" or donnees[:5] == b"%PDF-":
            texte, total = _texte_pdf(donnees, pages)
            meta["page_count"] = total
        elif ext in (".docx", ".dotx"):
            texte = _texte_docx(donnees)
        elif ext in (".html", ".htm"):
            texte = _texte_html(donnees)
        elif ext in EXT_TEXTE or not ext:
            texte = donnees.decode("utf-8", "replace")
        else:
            return meta | {"status": "unsupported", "detail": f"No extractor for {ext or 'this file'}."}
    except ImportError as err:
        return meta | {"status": "missing_extra", "detail": f"{err}. Install moodle-watch[documents]."}
    except Exception as err:  # noqa: BLE001
        return meta | {"status": "extract_failed", "detail": f"{type(err).__name__}: {err}"}

    tronque = len(texte) > max_chars
    return meta | {
        "status": "ok",
        "source": source,
        "chars": len(texte),
        "truncated": tronque,
        "text": texte[:max_chars],
    }
